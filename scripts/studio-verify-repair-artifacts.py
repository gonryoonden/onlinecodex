#!/usr/bin/env python3
"""Reopen actual repair browser downloads; native objects are required."""
from pathlib import Path
import json,zipfile
from PIL import Image
from pptx import Presentation
import fitz
root=Path(__file__).resolve().parents[1]/'evidence/studio/phase3'
report={}
for stem in ['local-repair','real-manual']:
 Image.open(root/(stem+'.png')).verify()
 pdf=fitz.open(root/(stem+'.pdf'));assert len(pdf)==2
 ppt=root/(stem+'.pptx');z=zipfile.ZipFile(ppt);assert z.testzip() is None
 slides=Presentation(ppt).slides;assert len(slides)==2
 texts=[s.text for slide in slides for s in slide.shapes if s.has_text_frame]
 expected='복원 후 한글 수정' if stem=='local-repair' else '2024년 교육 안내 · 직접 수정됨'
 assert any(expected in t for t in texts),texts
 assert any(s.shape_type==13 for s in slides[1].shapes)
 report[stem]={'pngReadable':True,'pdfPages':len(pdf),'pptxSlides':len(slides),'nativeText':texts,'pictures':sum(s.shape_type==13 for s in slides[1].shapes)}
(root/'repair-artifacts.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(report,ensure_ascii=False,indent=2))
