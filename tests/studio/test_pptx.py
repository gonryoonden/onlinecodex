import base64,copy,importlib.util,io,json,tempfile,unittest,zipfile
from pathlib import Path
from lxml import etree
from PIL import Image
ROOT=Path(__file__).resolve().parents[2]
spec=importlib.util.spec_from_file_location('studio_pptx',ROOT/'scripts/studio-pptx.py');adapter=importlib.util.module_from_spec(spec);spec.loader.exec_module(adapter)
NS=adapter.NS
class StudioPptxTests(unittest.TestCase):
 def setUp(self):self.project=json.loads((ROOT/'tests/studio/fixtures/project.json').read_text())
 def test_two_native_pages_and_source_alpha(self):
  with tempfile.TemporaryDirectory() as d:
   data=adapter.build(self.project,Path(d));z=zipfile.ZipFile(io.BytesIO(data));self.assertIsNone(z.testzip())
   slides=[n for n in z.namelist() if __import__('re').fullmatch('ppt/slides/slide[0-9]+.xml',n)];self.assertEqual(len(slides),2)
   first=etree.fromstring(z.read(slides[0]));second=etree.fromstring(z.read(slides[1]));self.assertIn('첫 페이지 한글 편집 가능',''.join(first.xpath('//a:t/text()',namespaces=NS)));self.assertIn('두 번째 페이지',''.join(second.xpath('//a:t/text()',namespaces=NS)))
   self.assertEqual(len(first.xpath('//p:pic',namespaces=NS)),1);self.assertGreaterEqual(len(first.xpath('//p:sp',namespaces=NS)),3)
   self.assertTrue(second.xpath('//a:gradFill',namespaces=NS));self.assertTrue(second.xpath('//a:xfrm[@rot="720000"]',namespaces=NS));self.assertTrue(second.xpath('//a:alpha[@val="70000"]',namespaces=NS))
   media=[n for n in z.namelist() if n.startswith('ppt/media/')];im=Image.open(io.BytesIO(z.read(media[0])));self.assertEqual(im.mode,'RGBA');self.assertEqual(im.getpixel((0,0))[3],0);self.assertEqual(im.getpixel((110,85))[3],255)
 def test_polygon_mask_certification_uses_unchanged_fidelity_validator(self):
  original=Image.open(io.BytesIO(base64.b64decode(self.project['assets']['original']['dataUrl'].split(',')[1]))).convert('RGBA')
  crop=original.crop((598,368,822,542));mask=Image.new('L',crop.size);from PIL import ImageDraw
  ImageDraw.Draw(mask).polygon([(112,2),(222,87),(112,172),(2,87)],fill=255);crop.putalpha(mask);b=io.BytesIO();crop.save(b,format='PNG')
  asset=self.project['assets']['diamond'];asset.update(width=224,height=174,dataUrl='data:image/png;base64,'+base64.b64encode(b.getvalue()).decode())
  element=self.project['pages'][0]['elements'][2];element.update(src=asset['dataUrl'],width=224,height=174,naturalWidth=224,naturalHeight=174);element['provenance']['sourcePixels']=[598,368,224,174]
  with tempfile.TemporaryDirectory() as d:
   adapter.build(self.project,Path(d));manifest=json.loads((Path(d)/'page-1.json').read_text());prov=manifest['asset_provenance'][0]
   self.assertEqual(prov['source_type'],'source-masked');self.assertEqual(adapter.validate_masked(prov['path'],prov['evidence']),[])
   image=Image.open(prov['path']);image.putpixel((112,87),(255,0,0,255));image.save(prov['path']);self.assertTrue(adapter.validate_masked(prov['path'],prov['evidence']))
 def test_exif_jpeg_orientation_matches_browser_coordinates(self):
  raw=(ROOT/'tests/studio/fixtures/synthetic-exif.jpg').read_bytes();a={'id':'exif','dataUrl':'data:image/jpeg;base64,'+base64.b64encode(raw).decode(),'mime':'image/jpeg','width':640,'height':960}
  with tempfile.TemporaryDirectory() as d:self.assertEqual(Image.open(adapter.decode_asset(a,Path(d))).size,(640,960))
 def test_all_xml_rgb_colors_are_six_digit_hex(self):
  with tempfile.TemporaryDirectory() as d:
   z=zipfile.ZipFile(io.BytesIO(adapter.build(self.project,Path(d))))
   for name in z.namelist():
    if name.endswith('.xml'):
     for color in etree.fromstring(z.read(name)).xpath('//a:srgbClr/@val',namespaces=NS):self.assertRegex(color,r'^[0-9A-F]{6}$')
 def test_single_page_opens_in_python_pptx(self):
  from pptx import Presentation
  with tempfile.TemporaryDirectory() as d:
   self.project['pages']=self.project['pages'][:1];data=adapter.build(self.project,Path(d));deck=Presentation(io.BytesIO(data));self.assertEqual(len(deck.slides),1);self.assertEqual(len(deck.slides[0].shapes),4)
 def test_relocated_asset_paths_are_absolute(self):
  import os
  with tempfile.TemporaryDirectory(dir=ROOT/'evidence/studio') as d:
   relative=Path(d).relative_to(ROOT);data=adapter.build(self.project,relative);self.assertTrue(data.startswith(b'PK'))
 def test_tampered_source_and_traversal_rejected(self):
  with tempfile.TemporaryDirectory() as d:
   self.project['assets']['diamond']['id']='../../secret'
   with self.assertRaises(ValueError):adapter.build(self.project,Path(d))
 def test_missing_asset_rejected(self):
  with tempfile.TemporaryDirectory() as d:
   del self.project['assets']['diamond']
   with self.assertRaises(ValueError):adapter.build(self.project,Path(d))
 def test_hidden_objects_not_exported(self):
  self.project['pages'][0]['elements'][-1]['visible']=False
  with tempfile.TemporaryDirectory() as d:
   z=zipfile.ZipFile(io.BytesIO(adapter.build(self.project,Path(d))));self.assertNotIn('첫 페이지 한글 편집 가능',z.read('ppt/slides/slide1.xml').decode())
if __name__=='__main__':unittest.main()
